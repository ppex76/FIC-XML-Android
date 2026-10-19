import csv
from pathlib import Path
from tkinter import filedialog, messagebox

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

import Studio_Dinamica_XML_NonElettroniche as base
import Studio_Dinamica_XML_NonElettroniche_v10 as v10


class App(v10.App):
    """Versione definitiva del riepilogo fiscale: il totale di riferimento è l'IMPONIBILE.

    Regola di inclusione invariata rispetto alla v10 validata sui file Excel storici:
    - include fatture il cui numero/numerazione NON contiene 'EL'
    - esclude fatture il cui numero/numerazione contiene 'EL'
    - pagamento, incasso, Sistema TS, stato SDI ed e_invoice API non filtrano i documenti

    Nei riepiloghi NON viene più mostrato amount_gross come totale principale, per evitare
    confusione con l'imponibile. I bolli restano separati.
    """

    def _build_ui(self):
        super()._build_ui()
        self.title("Studio Dinamica — NON elettroniche + imponibile + bolli v11")
        self.status.set(
            "Pronto. V11: il totale di riferimento è l'IMPONIBILE. "
            "Le fatture con sigla EL sono escluse; pagamento e Sistema TS non filtrano nulla."
        )

        # Riepilogo mensile: soltanto i valori richiesti/validati.
        self._summary_labels = {
            "anno": "Anno",
            "mese": "Mese",
            "fatture": "N. fatture",
            "imponibile": "Imponibile €",
            "bolli_n": "Marche da 2 €",
            "bolli_tot": "Totale bolli €",
        }
        cols = tuple(self._summary_labels.keys())
        self.summary_tree.configure(columns=cols)
        widths = {
            "anno": (75, "center"),
            "mese": (145, "w"),
            "fatture": (105, "e"),
            "imponibile": (170, "e"),
            "bolli_n": (135, "e"),
            "bolli_tot": (155, "e"),
        }
        for col in cols:
            width, anchor = widths[col]
            self.summary_tree.column(col, width=width, anchor=anchor)
            self.summary_tree.heading(
                col,
                text=self._summary_labels[col],
                command=lambda c=col: self.sort_summary_by(c),
            )

        # Dettaglio: niente colonna "Totale documento" nel prospetto fiscale.
        self._invoice_labels = {
            "anno": "Anno",
            "mese": "Mese",
            "data": "Data",
            "numero": "N. fattura",
            "cliente": "Cliente",
            "imponibile": "Imponibile €",
            "bollo": "Bollo €",
            "xml": "XML / stato",
        }
        icols = tuple(self._invoice_labels.keys())
        self.invoice_tree.configure(columns=icols)
        iwidths = {
            "anno": (70, "center"),
            "mese": (105, "w"),
            "data": (105, "center"),
            "numero": (110, "w"),
            "cliente": (300, "w"),
            "imponibile": (125, "e"),
            "bollo": (90, "e"),
            "xml": (330, "w"),
        }
        for col in icols:
            width, anchor = iwidths[col]
            self.invoice_tree.column(col, width=width, anchor=anchor)
            self.invoice_tree.heading(
                col,
                text=self._invoice_labels[col],
                command=lambda c=col: self.sort_invoices_by(c),
            )

        self._update_summary_headers()
        self._update_invoice_headers()

    def load_year(self):
        def job():
            try:
                c, t = self.creds()
                start, end = self.year_bounds()
                all_rows = self._list_all_invoices(c, t, start, end)

                included = []
                excluded_el = []
                for d in all_rows:
                    if self._is_electronic_by_number(d):
                        excluded_el.append(d)
                        continue
                    d["_stamp_total"], d["_stamp_count"] = base.stamp_info(d)
                    d["_xml_status"] = "Da esportare"
                    included.append(d)

                included.sort(key=lambda x: (
                    str(x.get("date") or ""),
                    base.invoice_number(x),
                    str(x.get("id") or ""),
                ))
                self.all_invoices = included
                self.after(0, self.render_all)

                jan = [d for d in included if str(d.get("date") or "")[5:7] == "01"]
                jan_net = sum(base.to_float(d.get("amount_net"), 0.0) for d in jan)
                self.set_status(
                    f"Caricate {len(included)} fatture NON elettroniche; escluse {len(excluded_el)} con sigla EL. "
                    f"GENNAIO: {len(jan)} fatture | IMPONIBILE € {base.fmt_money(jan_net)}. "
                    "Pagamento e Sistema TS non sono filtri."
                )
            except Exception as e:
                self.set_status("Errore: " + str(e))

        self.run_bg(job)

    def render_summary(self):
        for iid in self.summary_tree.get_children():
            self.summary_tree.delete(iid)
        year = self.year.get()
        for m in range(1, 13):
            r = self.summary_rows[m]
            self.summary_tree.insert(
                "", "end", iid=str(m),
                values=(
                    year,
                    base.MONTHS[m - 1],
                    r["count"],
                    base.fmt_money(r["net"]),
                    r["stamp_count"],
                    base.fmt_money(r["stamp_total"]),
                ),
            )
        self._sort_tree(self.summary_tree, self._summary_sort_col, self._summary_sort_reverse)
        self._update_summary_headers()

    def render_invoices(self):
        for iid in self.invoice_tree.get_children():
            self.invoice_tree.delete(iid)

        if self.selected_month:
            key = f"{self.selected_month:02d}"
            rows = [d for d in self.all_invoices if str(d.get("date") or "")[5:7] == key]
        else:
            rows = list(self.all_invoices)

        self.visible_invoices = rows
        for i, d in enumerate(rows):
            date_s = str(d.get("date") or "")
            try:
                year_s = date_s[:4]
                month_n = int(date_s[5:7])
                month_s = base.MONTHS[month_n - 1] if 1 <= month_n <= 12 else ""
            except Exception:
                year_s = ""
                month_s = ""
            ent = d.get("entity") or {}
            self.invoice_tree.insert(
                "", "end", iid=str(i),
                values=(
                    year_s,
                    month_s,
                    date_s,
                    base.invoice_number(d),
                    str(ent.get("name") or ""),
                    base.fmt_money(d.get("amount_net")),
                    base.fmt_money(d.get("_stamp_total")),
                    str(d.get("_xml_status") or ""),
                ),
            )

        self._sort_tree(self.invoice_tree, self._invoice_sort_col, self._invoice_sort_reverse)
        self._update_invoice_headers()

    def update_annual_line(self):
        net = sum(base.to_float(d.get("amount_net"), 0.0) for d in self.all_invoices)
        stamp_total = sum(base.to_float(d.get("_stamp_total"), 0.0) for d in self.all_invoices)
        stamp_count = sum(int(d.get("_stamp_count") or 0) for d in self.all_invoices)
        self.annual.set(
            f"ANNO {self.year.get()} — Fatture NON elettroniche: {len(self.all_invoices)} | "
            f"IMPONIBILE: € {base.fmt_money(net)} | Marche da 2 €: {stamp_count} | "
            f"Totale bolli: € {base.fmt_money(stamp_total)}"
        )

    def _write_detail_sheet(self, wb, docs, title="Fatture"):
        wf = wb.create_sheet(title)
        wf.freeze_panes = "A2"
        headers = [
            "Data", "Numero fattura", "Cliente", "Imponibile €",
            "Bollo €", "Marche da 2 €", "ID Fatture in Cloud"
        ]
        wf.append(headers)
        self._style_header(wf, 1)

        for d in docs:
            ent = d.get("entity") or {}
            wf.append([
                str(d.get("date") or ""),
                base.invoice_number(d),
                str(ent.get("name") or ""),
                base.to_float(d.get("amount_net"), 0.0),
                base.to_float(d.get("_stamp_total"), 0.0),
                int(d.get("_stamp_count") or 0),
                str(d.get("id") or ""),
            ])

        for rr in range(2, wf.max_row + 1):
            wf.cell(rr, 4).number_format = '#,##0.00 [$€-it-IT]'
            wf.cell(rr, 5).number_format = '#,##0.00 [$€-it-IT]'

        if wf.max_row >= 2:
            table = Table(
                displayName=f"TblV11{abs(hash((title, self.year.get(), len(docs)))) % 100000000}",
                ref=f"A1:G{wf.max_row}",
            )
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
                showRowStripes=True, showColumnStripes=False,
            )
            wf.add_table(table)

        for i, width in enumerate([12, 17, 38, 18, 13, 16, 20], 1):
            wf.column_dimensions[get_column_letter(i)].width = width

        if docs:
            footer = wf.max_row + 2
            wf.cell(footer, 3, "TOTALI")
            wf.cell(footer, 4, f"=SUM(D2:D{footer - 2})")
            wf.cell(footer, 5, f"=SUM(E2:E{footer - 2})")
            wf.cell(footer, 6, f"=SUM(F2:F{footer - 2})")
            for cc in range(3, 7):
                wf.cell(footer, cc).font = Font(bold=True)
                wf.cell(footer, cc).fill = PatternFill("solid", fgColor="D9EAF7")
            wf.cell(footer, 4).number_format = '#,##0.00 [$€-it-IT]'
            wf.cell(footer, 5).number_format = '#,##0.00 [$€-it-IT]'
        return wf

    def write_month_excel(self, path):
        month = self.month_number()
        docs = self.month_docs()
        year = int(self.year.get())
        company = self.company_name or "Studio Dinamica"
        r = self.summary_rows.get(month, {})

        wb = Workbook()
        ws = wb.active
        ws.title = "Riepilogo mese"
        ws.merge_cells("A1:E1")
        ws["A1"] = f"Studio Dinamica — {base.MONTHS[month - 1]} {year} — fatture NON elettroniche"
        self._style_title(ws["A1"])
        ws["A3"] = "Studio/Azienda"; ws["B3"] = company
        ws["A4"] = "Criterio"; ws["B4"] = "Numero/numerazione fattura NON contiene EL"
        headers = ["Mese", "N. fatture", "Imponibile €", "N. marche da 2 €", "Totale bolli €"]
        for cc, h in enumerate(headers, 1): ws.cell(6, cc, h)
        self._style_header(ws, 6)
        vals = [
            base.MONTHS[month - 1], int(r.get("count") or 0), base.to_float(r.get("net"), 0.0),
            int(r.get("stamp_count") or 0), base.to_float(r.get("stamp_total"), 0.0)
        ]
        for cc, v in enumerate(vals, 1): ws.cell(7, cc, v)
        for cc in (3, 5): ws.cell(7, cc).number_format = '#,##0.00 [$€-it-IT]'
        for i, w in enumerate([20, 14, 18, 19, 18], 1): ws.column_dimensions[get_column_letter(i)].width = w
        self._write_detail_sheet(wb, docs)
        path = Path(path); path.parent.mkdir(parents=True, exist_ok=True); wb.save(path)

    def export_month_excel(self):
        if not self.summary_rows:
            messagebox.showinfo(base.APP_TITLE, "Carica prima le fatture dell'anno.")
            return
        month = self.month_number()
        p = filedialog.asksaveasfilename(
            title="Esporta mese in Excel", defaultextension=".xlsx",
            initialfile=f"Studio_Dinamica_{self.year.get()}_{month:02d}_{base.MONTHS[month-1]}.xlsx",
            filetypes=[("Excel", "*.xlsx")],
        )
        if not p: return
        try:
            self.write_month_excel(Path(p))
            docs = self.month_docs()
            net = sum(base.to_float(d.get("amount_net"), 0.0) for d in docs)
            stamp = sum(base.to_float(d.get("_stamp_total"), 0.0) for d in docs)
            self.set_status(
                f"Excel {base.MONTHS[month-1]} salvato: {p} — {len(docs)} fatture, "
                f"IMPONIBILE € {base.fmt_money(net)}, bolli € {base.fmt_money(stamp)}"
            )
        except Exception as e:
            messagebox.showerror(base.APP_TITLE, "Errore durante l'esportazione Excel del mese:\n\n" + str(e))

    def write_quarter_excel(self, path):
        start_m, end_m, qcode = self.quarter_spec()
        docs = self.quarter_docs()
        year = int(self.year.get())
        wb = Workbook(); ws = wb.active; ws.title = "Riepilogo trimestre"
        ws.merge_cells("A1:E1")
        ws["A1"] = f"Studio Dinamica — {qcode} {year} — fatture NON elettroniche"
        self._style_title(ws["A1"])
        ws["A3"] = "Periodo"; ws["B3"] = f"{base.MONTHS[start_m-1]} - {base.MONTHS[end_m-1]} {year}"
        ws["A4"] = "Criterio"; ws["B4"] = "Numero/numerazione fattura NON contiene EL"
        headers = ["Mese", "N. fatture", "Imponibile €", "N. marche da 2 €", "Totale bolli €"]
        for cc, h in enumerate(headers, 1): ws.cell(6, cc, h)
        self._style_header(ws, 6)
        row = 7
        for m in range(start_m, end_m + 1):
            r = self.summary_rows.get(m, {})
            vals = [
                base.MONTHS[m-1], int(r.get("count") or 0), base.to_float(r.get("net"), 0.0),
                int(r.get("stamp_count") or 0), base.to_float(r.get("stamp_total"), 0.0)
            ]
            for cc, v in enumerate(vals, 1): ws.cell(row, cc, v)
            for cc in (3, 5): ws.cell(row, cc).number_format = '#,##0.00 [$€-it-IT]'
            row += 1
        tr = row
        ws.cell(tr, 1, "TOTALE TRIMESTRE")
        for cc in range(2, 6):
            ws.cell(tr, cc, f"=SUM({get_column_letter(cc)}7:{get_column_letter(cc)}{tr-1})")
        for cc in range(1, 6):
            ws.cell(tr, cc).font = Font(bold=True)
            ws.cell(tr, cc).fill = PatternFill("solid", fgColor="D9EAF7")
        for cc in (3, 5): ws.cell(tr, cc).number_format = '#,##0.00 [$€-it-IT]'
        for i, w in enumerate([20, 14, 18, 19, 18], 1): ws.column_dimensions[get_column_letter(i)].width = w
        self._write_detail_sheet(wb, docs)
        path = Path(path); path.parent.mkdir(parents=True, exist_ok=True); wb.save(path)

    def export_quarter_excel(self):
        if not self.summary_rows:
            messagebox.showinfo(base.APP_TITLE, "Carica prima le fatture dell'anno.")
            return
        start_m, end_m, qcode = self.quarter_spec()
        p = filedialog.asksaveasfilename(
            title="Esporta trimestre in Excel", defaultextension=".xlsx",
            initialfile=f"Studio_Dinamica_{self.year.get()}_{qcode}_{base.MONTHS[start_m-1]}-{base.MONTHS[end_m-1]}.xlsx",
            filetypes=[("Excel", "*.xlsx")],
        )
        if not p: return
        try:
            self.write_quarter_excel(Path(p))
            docs = self.quarter_docs()
            net = sum(base.to_float(d.get("amount_net"), 0.0) for d in docs)
            stamp = sum(base.to_float(d.get("_stamp_total"), 0.0) for d in docs)
            self.set_status(
                f"Excel {qcode} salvato: {p} — {len(docs)} fatture, "
                f"IMPONIBILE € {base.fmt_money(net)}, bolli € {base.fmt_money(stamp)}"
            )
        except Exception as e:
            messagebox.showerror(base.APP_TITLE, "Errore durante l'esportazione Excel del trimestre:\n\n" + str(e))

    def write_summary_csv(self, path):
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        total_net = 0.0
        total_stamp = 0.0
        total_stamp_count = 0
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f, delimiter=";")
            w.writerow(["Studio/Azienda", self.company_name or "Studio Dinamica"])
            w.writerow(["Anno", self.year.get()])
            w.writerow(["Criterio", "Numero/numerazione fattura NON contiene EL"])
            w.writerow([])
            w.writerow(["Mese", "N. fatture non elettroniche", "Imponibile", "N. marche da 2 euro", "Totale bolli"])
            for m in range(1, 13):
                r = self.summary_rows.get(m, {})
                net = base.to_float(r.get("net"), 0.0)
                stamp = base.to_float(r.get("stamp_total"), 0.0)
                stamp_count = int(r.get("stamp_count") or 0)
                total_net += net; total_stamp += stamp; total_stamp_count += stamp_count
                w.writerow([base.MONTHS[m-1], int(r.get("count") or 0), base.fmt_money(net), stamp_count, base.fmt_money(stamp)])
            w.writerow([])
            w.writerow(["TOTALE ANNO", len(self.all_invoices), base.fmt_money(total_net), total_stamp_count, base.fmt_money(total_stamp)])


if __name__ == "__main__":
    App().mainloop()
