import json
from datetime import date
from pathlib import Path
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

import Studio_Dinamica_XML_NonElettroniche as base
import Studio_Dinamica_XML_NonElettroniche_v2 as v2


class App(v2.App):
    def _build_ui(self):
        super()._build_ui()
        self.title("Studio Dinamica — XML + Excel mensile e trimestrale")
        self.geometry("1280x900")
        self.minsize(1050, 760)

        self.excel_month = tk.StringVar(value=base.MONTHS[date.today().month - 1])

        month_frame = ttk.Labelframe(self, text="Esportazione mensile Excel", padding=(12, 8))
        month_frame.grid(row=5, column=0, sticky="ew", padx=14, pady=(0, 12))
        month_frame.columnconfigure(3, weight=1)

        ttk.Label(month_frame, text="Mese:").grid(row=0, column=0, sticky="w")
        ttk.Combobox(
            month_frame,
            textvariable=self.excel_month,
            values=base.MONTHS,
            state="readonly",
            width=20,
        ).grid(row=0, column=1, padx=(8, 10), sticky="w")

        ttk.Button(
            month_frame,
            text="ESPORTA MESE IN EXCEL",
            command=self.export_month_excel,
        ).grid(row=0, column=2, padx=(0, 12), sticky="w")

        ttk.Label(
            month_frame,
            text="Excel del singolo mese con riepilogo + dettaglio di tutte le fatture non elettroniche.",
            foreground="#667085",
        ).grid(row=0, column=3, sticky="w")

    def _load_cfg(self):
        super()._load_cfg()
        try:
            d = json.loads(self.cfg_path.read_text("utf-8"))
            m = d.get("excel_month")
            if m in base.MONTHS:
                self.excel_month.set(m)
        except Exception:
            pass

    def _save_cfg(self):
        try:
            self.cfg_path.write_text(json.dumps({
                "company_id": self.company.get().strip(),
                "year": self.year.get(),
                "output_base": str(self.output_base) if self.output_base else "",
                "quarter": self.quarter.get() if hasattr(self, "quarter") else "",
                "excel_month": self.excel_month.get() if hasattr(self, "excel_month") else "",
            }, ensure_ascii=False, indent=2), encoding="utf-8")
        except Exception:
            pass

    def on_month_select(self, event=None):
        super().on_month_select(event)
        if self.selected_month:
            self.excel_month.set(base.MONTHS[self.selected_month - 1])

    def month_number(self):
        try:
            return base.MONTHS.index(self.excel_month.get()) + 1
        except ValueError:
            return 1

    def month_docs(self):
        m = self.month_number()
        key = f"{m:02d}"
        return [d for d in self.all_invoices if str(d.get("date") or "")[5:7] == key]

    @staticmethod
    def _style_title(cell):
        cell.font = Font(size=16, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    def write_month_excel(self, path):
        month = self.month_number()
        docs = self.month_docs()
        year = int(self.year.get())
        company = self.company_name or "Studio Dinamica"
        r = self.summary_rows.get(month, {})

        wb = Workbook()
        ws = wb.active
        ws.title = "Riepilogo mese"
        ws.freeze_panes = "A7"

        ws.merge_cells("A1:E1")
        ws["A1"] = f"Studio Dinamica — {base.MONTHS[month - 1]} {year}"
        self._style_title(ws["A1"])
        ws.row_dimensions[1].height = 25

        ws["A3"] = "Studio/Azienda"
        ws["B3"] = company
        ws["A4"] = "Periodo"
        ws["B4"] = f"{base.MONTHS[month - 1]} {year}"

        headers = [
            "Mese", "N. fatture non elettroniche", "Totale fatture €",
            "N. marche da 2 €", "Totale bolli €"
        ]
        for col, value in enumerate(headers, 1):
            ws.cell(row=6, column=col, value=value)
        self._style_header(ws, 6)

        ws.cell(row=7, column=1, value=base.MONTHS[month - 1])
        ws.cell(row=7, column=2, value=int(r.get("count") or 0))
        ws.cell(row=7, column=3, value=base.to_float(r.get("total"), 0.0))
        ws.cell(row=7, column=4, value=int(r.get("stamp_count") or 0))
        ws.cell(row=7, column=5, value=base.to_float(r.get("stamp_total"), 0.0))

        ws.cell(row=8, column=1, value="TOTALE MESE")
        ws.cell(row=8, column=2, value="=B7")
        ws.cell(row=8, column=3, value="=C7")
        ws.cell(row=8, column=4, value="=D7")
        ws.cell(row=8, column=5, value="=E7")
        for cell in ws[8]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")

        for rr in (7, 8):
            ws.cell(rr, 3).number_format = '#,##0.00 [$€-it-IT]'
            ws.cell(rr, 5).number_format = '#,##0.00 [$€-it-IT]'

        for col, width in {"A": 20, "B": 28, "C": 20, "D": 19, "E": 18}.items():
            ws.column_dimensions[col].width = width

        ws.cell(10, 1, "Nota")
        ws.cell(10, 2, "Totale fatture = amount_gross Fatture in Cloud. I bolli sono esposti separatamente per evitare doppi conteggi.")
        ws.merge_cells("B10:E10")
        ws["B10"].alignment = Alignment(wrap_text=True)

        wf = wb.create_sheet("Fatture")
        wf.freeze_panes = "A2"
        detail_headers = [
            "Data", "Numero fattura", "Cliente", "Totale fattura €",
            "Bollo €", "Marche da 2 €", "ID Fatture in Cloud"
        ]
        wf.append(detail_headers)
        self._style_header(wf, 1)

        for d in docs:
            ent = d.get("entity") or {}
            wf.append([
                str(d.get("date") or ""),
                base.invoice_number(d),
                str(ent.get("name") or ""),
                base.to_float(d.get("amount_gross"), 0.0),
                base.to_float(d.get("_stamp_total"), 0.0),
                int(d.get("_stamp_count") or 0),
                str(d.get("id") or ""),
            ])

        for rr in range(2, wf.max_row + 1):
            wf.cell(rr, 4).number_format = '#,##0.00 [$€-it-IT]'
            wf.cell(rr, 5).number_format = '#,##0.00 [$€-it-IT]'

        if wf.max_row >= 2:
            table_name = f"FattureM{month:02d}{year}"
            table = Table(displayName=table_name, ref=f"A1:G{wf.max_row}")
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            wf.add_table(table)

        for i, width in enumerate([12, 16, 40, 18, 14, 16, 20], 1):
            wf.column_dimensions[get_column_letter(i)].width = width

        if docs:
            footer = wf.max_row + 2
            wf.cell(footer, 3, "TOTALI MESE")
            wf.cell(footer, 4, f"=SUM(D2:D{footer - 2})")
            wf.cell(footer, 5, f"=SUM(E2:E{footer - 2})")
            wf.cell(footer, 6, f"=SUM(F2:F{footer - 2})")
            for c in range(3, 7):
                wf.cell(footer, c).font = Font(bold=True)
                wf.cell(footer, c).fill = PatternFill("solid", fgColor="D9EAF7")
            wf.cell(footer, 4).number_format = '#,##0.00 [$€-it-IT]'
            wf.cell(footer, 5).number_format = '#,##0.00 [$€-it-IT]'

        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)

    def export_month_excel(self):
        if not self.summary_rows:
            messagebox.showinfo(base.APP_TITLE, "Carica prima le fatture dell'anno.")
            return

        month = self.month_number()
        self._save_cfg()
        suggested = f"Studio_Dinamica_{self.year.get()}_{month:02d}_{base.MONTHS[month - 1]}.xlsx"
        p = filedialog.asksaveasfilename(
            title="Esporta mese in Excel",
            defaultextension=".xlsx",
            initialfile=suggested,
            filetypes=[("Excel", "*.xlsx")],
        )
        if not p:
            return

        try:
            self.write_month_excel(Path(p))
            docs = self.month_docs()
            total = sum(base.to_float(d.get("amount_gross"), 0.0) for d in docs)
            stamp_total = sum(base.to_float(d.get("_stamp_total"), 0.0) for d in docs)
            self.set_status(
                f"Excel {base.MONTHS[month - 1]} salvato: {p} — {len(docs)} fatture, "
                f"totale € {base.fmt_money(total)}, bolli € {base.fmt_money(stamp_total)}"
            )
        except Exception as e:
            messagebox.showerror(base.APP_TITLE, "Errore durante l'esportazione Excel del mese:\n\n" + str(e))


if __name__ == "__main__":
    App().mainloop()
