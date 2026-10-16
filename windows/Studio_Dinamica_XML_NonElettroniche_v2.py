import json
from datetime import date
from pathlib import Path
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

import Studio_Dinamica_XML_NonElettroniche as base


QUARTERS = {
    "1° trimestre — Gennaio / Marzo": (1, 3, "Q1"),
    "2° trimestre — Aprile / Giugno": (4, 6, "Q2"),
    "3° trimestre — Luglio / Settembre": (7, 9, "Q3"),
    "4° trimestre — Ottobre / Dicembre": (10, 12, "Q4"),
}


class App(base.App):
    def _build_ui(self):
        super()._build_ui()
        self.title("Studio Dinamica — XML + Excel trimestrale")
        q_now = ((date.today().month - 1) // 3) + 1
        self.quarter = tk.StringVar(value=list(QUARTERS.keys())[q_now - 1])

        quarter_frame = ttk.Labelframe(self, text="Esportazione trimestrale Excel", padding=(12, 8))
        quarter_frame.grid(row=4, column=0, sticky="ew", padx=14, pady=(0, 12))
        quarter_frame.columnconfigure(3, weight=1)
        ttk.Label(quarter_frame, text="Trimestre:").grid(row=0, column=0, sticky="w")
        ttk.Combobox(
            quarter_frame,
            textvariable=self.quarter,
            values=list(QUARTERS.keys()),
            state="readonly",
            width=34,
        ).grid(row=0, column=1, padx=(8, 10), sticky="w")
        ttk.Button(
            quarter_frame,
            text="ESPORTA TRIMESTRE IN EXCEL",
            command=self.export_quarter_excel,
        ).grid(row=0, column=2, padx=(0, 12), sticky="w")
        ttk.Label(
            quarter_frame,
            text="Excel con riepilogo dei 3 mesi + dettaglio di tutte le fatture del trimestre.",
            foreground="#667085",
        ).grid(row=0, column=3, sticky="w")

    def _load_cfg(self):
        super()._load_cfg()
        try:
            d = json.loads(self.cfg_path.read_text("utf-8"))
            q = d.get("quarter")
            if q in QUARTERS:
                self.quarter.set(q)
        except Exception:
            pass

    def _save_cfg(self):
        try:
            self.cfg_path.write_text(json.dumps({
                "company_id": self.company.get().strip(),
                "year": self.year.get(),
                "output_base": str(self.output_base) if self.output_base else "",
                "quarter": self.quarter.get() if hasattr(self, "quarter") else "",
            }, ensure_ascii=False, indent=2), encoding="utf-8")
        except Exception:
            pass

    def on_month_select(self, event=None):
        super().on_month_select(event)
        if self.selected_month:
            q_index = (self.selected_month - 1) // 3
            self.quarter.set(list(QUARTERS.keys())[q_index])

    def quarter_spec(self):
        return QUARTERS.get(self.quarter.get(), (1, 3, "Q1"))

    def quarter_docs(self):
        start_m, end_m, _ = self.quarter_spec()
        docs = []
        for d in self.all_invoices:
            try:
                m = int(str(d.get("date") or "")[5:7])
            except Exception:
                continue
            if start_m <= m <= end_m:
                docs.append(d)
        return docs

    @staticmethod
    def _style_header(ws, row):
        fill = PatternFill("solid", fgColor="1F4E78")
        font = Font(color="FFFFFF", bold=True)
        side = Side(style="thin", color="D9E1F2")
        for cell in ws[row]:
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=side)

    def write_quarter_excel(self, path):
        start_m, end_m, qcode = self.quarter_spec()
        docs = self.quarter_docs()
        year = int(self.year.get())
        company = self.company_name or "Studio Dinamica"

        wb = Workbook()
        ws = wb.active
        ws.title = "Riepilogo trimestre"
        ws.freeze_panes = "A7"

        ws.merge_cells("A1:E1")
        ws["A1"] = f"Studio Dinamica — {qcode} {year}"
        ws["A1"].font = Font(size=16, bold=True, color="FFFFFF")
        ws["A1"].fill = PatternFill("solid", fgColor="1F4E78")
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 25

        ws["A3"] = "Studio/Azienda"
        ws["B3"] = company
        ws["A4"] = "Periodo"
        ws["B4"] = f"{base.MONTHS[start_m - 1]} - {base.MONTHS[end_m - 1]} {year}"
        ws["D3"] = "Trimestre"
        ws["E3"] = qcode

        headers = [
            "Mese", "N. fatture non elettroniche", "Totale fatture €",
            "N. marche da 2 €", "Totale bolli €"
        ]
        for col, value in enumerate(headers, 1):
            ws.cell(row=6, column=col, value=value)
        self._style_header(ws, 6)

        row = 7
        for m in range(start_m, end_m + 1):
            r = self.summary_rows.get(m, {})
            ws.cell(row=row, column=1, value=base.MONTHS[m - 1])
            ws.cell(row=row, column=2, value=int(r.get("count") or 0))
            ws.cell(row=row, column=3, value=base.to_float(r.get("total"), 0.0))
            ws.cell(row=row, column=4, value=int(r.get("stamp_count") or 0))
            ws.cell(row=row, column=5, value=base.to_float(r.get("stamp_total"), 0.0))
            row += 1

        total_row = row
        ws.cell(total_row, 1, "TOTALE TRIMESTRE")
        ws.cell(total_row, 2, f"=SUM(B7:B{total_row - 1})")
        ws.cell(total_row, 3, f"=SUM(C7:C{total_row - 1})")
        ws.cell(total_row, 4, f"=SUM(D7:D{total_row - 1})")
        ws.cell(total_row, 5, f"=SUM(E7:E{total_row - 1})")
        for cell in ws[total_row]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")

        for r in range(7, total_row + 1):
            ws.cell(r, 3).number_format = '#,##0.00 [$€-it-IT]'
            ws.cell(r, 5).number_format = '#,##0.00 [$€-it-IT]'

        for col, width in {"A": 20, "B": 28, "C": 20, "D": 19, "E": 18}.items():
            ws.column_dimensions[col].width = width

        note_row = total_row + 2
        ws.cell(note_row, 1, "Nota")
        ws.cell(
            note_row, 2,
            "Totale fatture = amount_gross Fatture in Cloud. I bolli sono esposti separatamente per evitare doppi conteggi."
        )
        ws.merge_cells(start_row=note_row, start_column=2, end_row=note_row, end_column=5)
        ws.cell(note_row, 2).alignment = Alignment(wrap_text=True)

        wf = wb.create_sheet("Fatture")
        wf.freeze_panes = "A2"
        detail_headers = [
            "Data", "Numero fattura", "Cliente", "Totale fattura €",
            "Bollo €", "Marche da 2 €", "ID Fatture in Cloud"
        ]
        wf.append(detail_headers)
        self._style_header(wf, 1)

        for d in docs:
            ent = d.get("entity") or {}
            wf.append([
                str(d.get("date") or ""),
                base.invoice_number(d),
                str(ent.get("name") or ""),
                base.to_float(d.get("amount_gross"), 0.0),
                base.to_float(d.get("_stamp_total"), 0.0),
                int(d.get("_stamp_count") or 0),
                str(d.get("id") or ""),
            ])

        for r in range(2, wf.max_row + 1):
            wf.cell(r, 4).number_format = '#,##0.00 [$€-it-IT]'
            wf.cell(r, 5).number_format = '#,##0.00 [$€-it-IT]'

        if wf.max_row >= 2:
            table = Table(displayName=f"Fatture{qcode}{year}", ref=f"A1:G{wf.max_row}")
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            wf.add_table(table)

        for i, width in enumerate([12, 16, 40, 18, 14, 16, 20], 1):
            wf.column_dimensions[get_column_letter(i)].width = width

        if docs:
            footer = wf.max_row + 2
            wf.cell(footer, 3, "TOTALI TRIMESTRE")
            wf.cell(footer, 4, f"=SUM(D2:D{footer - 2})")
            wf.cell(footer, 5, f"=SUM(E2:E{footer - 2})")
            wf.cell(footer, 6, f"=SUM(F2:F{footer - 2})")
            for c in range(3, 7):
                wf.cell(footer, c).font = Font(bold=True)
                wf.cell(footer, c).fill = PatternFill("solid", fgColor="D9EAF7")
            wf.cell(footer, 4).number_format = '#,##0.00 [$€-it-IT]'
            wf.cell(footer, 5).number_format = '#,##0.00 [$€-it-IT]'

        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)

    def export_quarter_excel(self):
        if not self.summary_rows:
            messagebox.showinfo(base.APP_TITLE, "Carica prima le fatture dell'anno.")
            return

        start_m, end_m, qcode = self.quarter_spec()
        self._save_cfg()
        suggested = (
            f"Studio_Dinamica_{self.year.get()}_{qcode}_"
            f"{base.MONTHS[start_m - 1]}-{base.MONTHS[end_m - 1]}.xlsx"
        )
        p = filedialog.asksaveasfilename(
            title="Esporta trimestre in Excel",
            defaultextension=".xlsx",
            initialfile=suggested,
            filetypes=[("Excel", "*.xlsx")],
        )
        if not p:
            return

        try:
            self.write_quarter_excel(Path(p))
            docs = self.quarter_docs()
            total = sum(base.to_float(d.get("amount_gross"), 0.0) for d in docs)
            stamp_total = sum(base.to_float(d.get("_stamp_total"), 0.0) for d in docs)
            self.set_status(
                f"Excel {qcode} salvato: {p} — {len(docs)} fatture, "
                f"totale € {base.fmt_money(total)}, bolli € {base.fmt_money(stamp_total)}"
            )
        except Exception as e:
            messagebox.showerror(base.APP_TITLE, "Errore durante l'esportazione Excel:\n\n" + str(e))


if __name__ == "__main__":
    App().mainloop()
