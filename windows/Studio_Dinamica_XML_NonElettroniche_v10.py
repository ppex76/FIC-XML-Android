from pathlib import Path
from tkinter import filedialog, messagebox

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

import Studio_Dinamica_XML_NonElettroniche as base
import Studio_Dinamica_XML_NonElettroniche_v9 as v9


class App(v9.App):
    def _build_ui(self):
        super()._build_ui()
        self.title("Studio Dinamica — NON elettroniche (regola EL) + Excel imponibile v10")
        self.status.set(
            "Pronto. V10: stessa regola del riepilogo validato: una fattura è elettronica solo se il numero/numerazione contiene 'EL'. "
            "Pagamento, TS, incasso ed e_invoice API non sono filtri."
        )
        # Il flag API resta visibile solo come diagnostica: NON decide l'inclusione.
        if hasattr(self, "_invoice_labels") and "elettronica" in self._invoice_labels:
            self._invoice_labels["elettronica"] = "Flag API e_invoice"
            self._update_invoice_headers()

    @staticmethod
    def _is_electronic_by_number(d):
        """Replica ESATTAMENTE la regola usata nel vecchio estrattore Excel validato.

        Il vecchio programma considerava elettronica una fattura se nel numero fattura
        compariva la stringa 'EL'. Qui applichiamo la stessa regola a number+numeration.
        """
        label = base.invoice_number(d)
        return "EL" in str(label or "").upper()

    def load_year(self):
        """Carica tutte le fatture e scarta solo quelle il cui numero/numerazione contiene EL.

        NON filtra per:
        - e_invoice API
        - stato SDI
        - pagamento/incasso
        - Sistema Tessera Sanitaria
        """
        def job():
            try:
                c, t = self.creds()
                start, end = self.year_bounds()
                all_rows = self._list_all_invoices(c, t, start, end)

                included = []
                excluded_el = []
                kept_api_einvoice = []

                for d in all_rows:
                    if self._is_electronic_by_number(d):
                        excluded_el.append(d)
                        continue

                    # e_invoice NON viene usato come filtro: alcune fatture del riepilogo
                    # validato possono avere e_invoice=True pur non avendo sigla EL.
                    if base.boolish(d.get("e_invoice")):
                        kept_api_einvoice.append(d)

                    d["_stamp_total"], d["_stamp_count"] = base.stamp_info(d)
                    d["_xml_status"] = "Da esportare"
                    included.append(d)

                included.sort(
                    key=lambda x: (
                        str(x.get("date") or ""),
                        base.invoice_number(x),
                        str(x.get("id") or ""),
                    )
                )
                self.all_invoices = included
                self.after(0, self.render_all)

                jan = [d for d in included if str(d.get("date") or "")[5:7] == "01"]
                jan_net = sum(base.to_float(d.get("amount_net"), 0.0) for d in jan)
                jan_gross = sum(base.to_float(d.get("amount_gross"), 0.0) for d in jan)
                jan_kept_api = [d for d in kept_api_einvoice if str(d.get("date") or "")[5:7] == "01"]
                jan_kept_api_net = sum(base.to_float(d.get("amount_net"), 0.0) for d in jan_kept_api)

                self.set_status(
                    f"V10 — {len(included)} fatture NON elettroniche secondo regola numero senza 'EL'; "
                    f"escluse {len(excluded_el)} con sigla EL. "
                    f"GENNAIO: {len(jan)} fatture | imponibile € {base.fmt_money(jan_net)} | "
                    f"totale documenti € {base.fmt_money(jan_gross)}. "
                    f"Incluse nonostante e_invoice=SI: {len(jan_kept_api)} a gennaio, imponibile € {base.fmt_money(jan_kept_api_net)}."
                )
            except Exception as e:
                self.set_status("Errore: " + str(e))

        self.run_bg(job)

    def update_annual_line(self):
        net = sum(base.to_float(d.get("amount_net"), 0.0) for d in self.all_invoices)
        gross = sum(base.to_float(d.get("amount_gross"), 0.0) for d in self.all_invoices)
        stamp_total = sum(base.to_float(d.get("_stamp_total"), 0.0) for d in self.all_invoices)
        stamp_count = sum(int(d.get("_stamp_count") or 0) for d in self.all_invoices)
        self.annual.set(
            f"ANNO {self.year.get()} — NON elettroniche (senza sigla EL): {len(self.all_invoices)} | "
            f"Imponibile: € {base.fmt_money(net)} | Totale documenti: € {base.fmt_money(gross)} | "
            f"Marche da 2 €: {stamp_count} | Totale bolli: € {base.fmt_money(stamp_total)}"
        )

    def _write_detail_sheet(self, wb, docs, title="Fatture"):
        wf = wb.create_sheet(title)
        wf.freeze_panes = "A2"
        headers = [
            "Data", "Numero fattura", "Cliente", "Imponibile €", "Totale documento €",
            "Bollo €", "Marche da 2 €", "Flag API e_invoice", "Stato SDI", "ID Fatture in Cloud"
        ]
        wf.append(headers)
        self._style_header(wf, 1)

        for d in docs:
            ent = d.get("entity") or {}
            wf.append([
                str(d.get("date") or ""),
                base.invoice_number(d),
                str(ent.get("name") or ""),
                base.to_float(d.get("amount_net"), 0.0),
                base.to_float(d.get("amount_gross"), 0.0),
                base.to_float(d.get("_stamp_total"), 0.0),
                int(d.get("_stamp_count") or 0),
                "SI" if base.boolish(d.get("e_invoice")) else "NO",
                str(d.get("ei_status") or ""),
                str(d.get("id") or ""),
            ])

        for rr in range(2, wf.max_row + 1):
            for cc in (4, 5, 6):
                wf.cell(rr, cc).number_format = '#,##0.00 [$€-it-IT]'

        if wf.max_row >= 2:
            table = Table(displayName=f"Tbl{abs(hash((title, self.year.get(), len(docs)))) % 100000000}", ref=f"A1:J{wf.max_row}")
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
                showRowStripes=True, showColumnStripes=False,
            )
            wf.add_table(table)

        widths = [12, 17, 38, 16, 19, 13, 16, 18, 16, 20]
        for i, width in enumerate(widths, 1):
            wf.column_dimensions[get_column_letter(i)].width = width

        if docs:
            footer = wf.max_row + 2
            wf.cell(footer, 3, "TOTALI")
            wf.cell(footer, 4, f"=SUM(D2:D{footer - 2})")
            wf.cell(footer, 5, f"=SUM(E2:E{footer - 2})")
            wf.cell(footer, 6, f"=SUM(F2:F{footer - 2})")
            wf.cell(footer, 7, f"=SUM(G2:G{footer - 2})")
            for c in range(3, 8):
                wf.cell(footer, c).font = Font(bold=True)
                wf.cell(footer, c).fill = PatternFill("solid", fgColor="D9EAF7")
            for c in (4, 5, 6):
                wf.cell(footer, c).number_format = '#,##0.00 [$€-it-IT]'
        return wf

    def write_month_excel(self, path):
        month = self.month_number()
        docs = self.month_docs()
        year = int(self.year.get())
        company = self.company_name or "Studio Dinamica"
        r = self.summary_rows.get(month, {})

        wb = Workbook()
        ws = wb.active
        ws.title = "Riepilogo mese"
        ws.merge_cells("A1:F1")
        ws["A1"] = f"Studio Dinamica — {base.MONTHS[month - 1]} {year} — fatture NON elettroniche"
        self._style_title(ws["A1"])
        ws["A3"] = "Studio/Azienda"; ws["B3"] = company
        ws["A4"] = "Criterio"; ws["B4"] = "Numero/numerazione fattura NON contiene EL"

        headers = ["Mese", "N. fatture", "Imponibile €", "Totale documenti €", "N. marche da 2 €", "Totale bolli €"]
        for c, h in enumerate(headers, 1): ws.cell(6, c, h)
        self._style_header(ws, 6)
        vals = [base.MONTHS[month - 1], int(r.get("count") or 0), base.to_float(r.get("net"), 0.0),
                base.to_float(r.get("total"), 0.0), int(r.get("stamp_count") or 0), base.to_float(r.get("stamp_total"), 0.0)]
        for c, v in enumerate(vals, 1): ws.cell(7, c, v)
        for c in (3,4,6): ws.cell(7,c).number_format = '#,##0.00 [$€-it-IT]'
        for i,w in enumerate([20,14,18,20,19,18],1): ws.column_dimensions[get_column_letter(i)].width=w

        self._write_detail_sheet(wb, docs)
        path = Path(path); path.parent.mkdir(parents=True, exist_ok=True); wb.save(path)

    def export_month_excel(self):
        if not self.summary_rows:
            messagebox.showinfo(base.APP_TITLE, "Carica prima le fatture dell'anno.")
            return
        month = self.month_number()
        p = filedialog.asksaveasfilename(
            title="Esporta mese in Excel", defaultextension=".xlsx",
            initialfile=f"Studio_Dinamica_{self.year.get()}_{month:02d}_{base.MONTHS[month-1]}.xlsx",
            filetypes=[("Excel", "*.xlsx")],
        )
        if not p: return
        try:
            self.write_month_excel(Path(p))
            docs = self.month_docs()
            net = sum(base.to_float(d.get("amount_net"),0.0) for d in docs)
            gross = sum(base.to_float(d.get("amount_gross"),0.0) for d in docs)
            self.set_status(f"Excel {base.MONTHS[month-1]} salvato: {p} — {len(docs)} fatture, imponibile € {base.fmt_money(net)}, totale documenti € {base.fmt_money(gross)}")
        except Exception as e:
            messagebox.showerror(base.APP_TITLE, "Errore durante l'esportazione Excel del mese:\n\n" + str(e))

    def write_quarter_excel(self, path):
        start_m, end_m, qcode = self.quarter_spec()
        docs = self.quarter_docs()
        year = int(self.year.get())
        wb = Workbook(); ws = wb.active; ws.title = "Riepilogo trimestre"
        ws.merge_cells("A1:F1"); ws["A1"] = f"Studio Dinamica — {qcode} {year} — fatture NON elettroniche"; self._style_title(ws["A1"])
        ws["A3"]="Periodo"; ws["B3"]=f"{base.MONTHS[start_m-1]} - {base.MONTHS[end_m-1]} {year}"
        ws["A4"]="Criterio"; ws["B4"]="Numero/numerazione fattura NON contiene EL"
        headers=["Mese","N. fatture","Imponibile €","Totale documenti €","N. marche da 2 €","Totale bolli €"]
        for c,h in enumerate(headers,1): ws.cell(6,c,h)
        self._style_header(ws,6)
        row=7
        for m in range(start_m,end_m+1):
            r=self.summary_rows.get(m,{})
            vals=[base.MONTHS[m-1],int(r.get("count") or 0),base.to_float(r.get("net"),0.0),base.to_float(r.get("total"),0.0),int(r.get("stamp_count") or 0),base.to_float(r.get("stamp_total"),0.0)]
            for c,v in enumerate(vals,1): ws.cell(row,c,v)
            for c in (3,4,6): ws.cell(row,c).number_format='#,##0.00 [$€-it-IT]'
            row+=1
        tr=row; ws.cell(tr,1,"TOTALE TRIMESTRE")
        for c in range(2,7): ws.cell(tr,c,f"=SUM({get_column_letter(c)}7:{get_column_letter(c)}{tr-1})")
        for c in range(1,7): ws.cell(tr,c).font=Font(bold=True); ws.cell(tr,c).fill=PatternFill("solid",fgColor="D9EAF7")
        for c in (3,4,6): ws.cell(tr,c).number_format='#,##0.00 [$€-it-IT]'
        for i,w in enumerate([20,14,18,20,19,18],1): ws.column_dimensions[get_column_letter(i)].width=w
        self._write_detail_sheet(wb, docs)
        path=Path(path); path.parent.mkdir(parents=True,exist_ok=True); wb.save(path)

    def export_quarter_excel(self):
        if not self.summary_rows:
            messagebox.showinfo(base.APP_TITLE, "Carica prima le fatture dell'anno.")
            return
        start_m,end_m,qcode=self.quarter_spec()
        p=filedialog.asksaveasfilename(title="Esporta trimestre in Excel",defaultextension=".xlsx",initialfile=f"Studio_Dinamica_{self.year.get()}_{qcode}_{base.MONTHS[start_m-1]}-{base.MONTHS[end_m-1]}.xlsx",filetypes=[("Excel","*.xlsx")])
        if not p: return
        try:
            self.write_quarter_excel(Path(p))
            docs=self.quarter_docs(); net=sum(base.to_float(d.get("amount_net"),0.0) for d in docs); gross=sum(base.to_float(d.get("amount_gross"),0.0) for d in docs)
            self.set_status(f"Excel {qcode} salvato: {p} — {len(docs)} fatture, imponibile € {base.fmt_money(net)}, totale documenti € {base.fmt_money(gross)}")
        except Exception as e:
            messagebox.showerror(base.APP_TITLE,"Errore durante l'esportazione Excel del trimestre:\n\n"+str(e))


if __name__ == "__main__":
    App().mainloop()
